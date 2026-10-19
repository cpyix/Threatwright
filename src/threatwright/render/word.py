"""DOCX: docxtpl renders Word templates whose body loops over the outline blocks.

Styles, header, footer and page setup come from the template file; the body
of the bundled templates (and of derived ones) is written by `write_body`.
"""

import datetime
from pathlib import Path
from typing import Any

from docx import Document as open_docx
from docx.document import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate  # type: ignore[import-untyped]
from jinja2.sandbox import ImmutableSandboxedEnvironment

# Block kind -> paragraph style used for its text.
PARAGRAPH_STYLES = {
    "title": "Title",
    "subtitle": "Subtitle",
    "h1": "Heading 1",
    "h2": "Heading 2",
    "p": "Normal",
    "note": "Note",
}


def render_docx(
    template: Path, ctx: dict[str, Any], env: ImmutableSandboxedEnvironment, out: Path
) -> None:
    tpl = DocxTemplate(template)
    tpl.render(ctx, jinja_env=env, autoescape=True)
    date = ctx["document"]["date"]
    if isinstance(date, datetime.date):
        stamp = datetime.datetime(date.year, date.month, date.day, tzinfo=datetime.UTC)
        tpl.docx.core_properties.created = stamp
        tpl.docx.core_properties.modified = stamp
    tpl.save(out)


def undeclared(template: Path, env: ImmutableSandboxedEnvironment) -> set[str]:
    return set(DocxTemplate(template).get_undeclared_template_variables(jinja_env=env))


# --- writing template bodies ------------------------------------------------------


def _shade(cell_xml: Any, fill: str) -> None:
    tc_pr = cell_xml.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set(cell_par: Paragraph, text: str, bold: bool = False) -> None:
    run = cell_par.add_run(text)
    run.bold = bold


def _table(doc: Document, rows: int, cols: int, style: str | None) -> Table:
    table: Table = doc.add_table(rows=rows, cols=cols)
    if style and style in [s.name for s in doc.styles]:
        table.style = doc.styles[style]
    return table


def ensure_styles(doc: Document) -> None:
    """Add the paragraph styles the body needs when a document lacks them."""
    names = {s.name for s in doc.styles}
    for name in PARAGRAPH_STYLES.values():
        if name in names:
            continue
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        font = style.font
        if name == "Title":
            font.size, font.bold = Pt(28), True
        elif name.startswith("Heading"):
            font.bold = True
            font.size = Pt(16 if name.endswith("1") else 13)
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.keep_with_next = True
        elif name == "Subtitle":
            font.size = Pt(14)
        elif name == "Note":
            font.italic = True
            style.paragraph_format.left_indent = Pt(18)


def write_body(
    doc: Document, table_style: str | None = "Table Grid", head_fill: str = "D9E2F3"
) -> None:
    """Append the generic outline loop (docxtpl tags) to the document body."""

    def par(text: str, style: str | None = None) -> Paragraph:
        p = doc.add_paragraph(text)
        if style:
            p.style = doc.styles[style]
        return p

    par("{%p for s in outline %}")
    par("{%p for b in s.blocks %}")
    for kind, style in PARAGRAPH_STYLES.items():
        par(f"{{%p if b.kind == '{kind}' %}}")
        par("{{ b.text }}", style)
        par("{%p endif %}")

    par("{%p if b.kind == 'pagebreak' %}")
    par("").add_run().add_break(WD_BREAK.PAGE)
    par("{%p endif %}")

    par("{%p if b.kind == 'table' %}")
    table = _table(doc, 4, 3, table_style)
    head = table.rows[0]
    tr_pr = head._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))
    for cell, text in zip(
        head.cells, ("{%tc for h in b.headers %}", "{{ h }}", "{%tc endfor %}"), strict=True
    ):
        _set(cell.paragraphs[0], text, bold=True)
        _shade(cell._tc, head_fill)
    _set(table.rows[1].cells[0].paragraphs[0], "{%tr for r in b.rows %}")
    for cell, text in zip(
        table.rows[2].cells, ("{%tc for v in r %}", "{{ v }}", "{%tc endfor %}"), strict=True
    ):
        _set(cell.paragraphs[0], text)
    _set(table.rows[3].cells[0].paragraphs[0], "{%tr endfor %}")
    par("")  # keeps consecutive tables apart (Word would join them)
    par("{%p endif %}")

    par("{%p if b.kind == 'fields' %}")
    fields = _table(doc, 3, 2, table_style)
    _set(fields.rows[0].cells[0].paragraphs[0], "{%tr for r in b.rows %}")
    _set(fields.rows[1].cells[0].paragraphs[0], "{{ r[0] }}", bold=True)
    _set(fields.rows[1].cells[1].paragraphs[0], "{{ r[1] }}")
    _shade(fields.rows[1].cells[0]._tc, head_fill)
    _set(fields.rows[2].cells[0].paragraphs[0], "{%tr endfor %}")
    par("")
    par("{%p endif %}")

    par("{%p endfor %}")
    par("{%p endfor %}")


def derive(existing: Path, target: Path) -> None:
    """Keep styles, header/footer and page setup; replace the body with the outline loop."""
    doc = open_docx(str(existing))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    ensure_styles(doc)
    write_body(doc)
    doc.save(str(target))
