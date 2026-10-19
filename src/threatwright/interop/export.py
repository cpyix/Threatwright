"""Threat and measure lists as CSV, JSON or XLSX.

CSV choices (German Excel opens it with a double click): UTF-8 *with* BOM, so
Excel detects the encoding, `;` as default delimiter (German list separator),
LF line endings. Cells starting with = + - @ TAB CR get a leading `'` against
formula injection. XLSX stores such cells as plain strings instead.
"""

import csv
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any, Literal

from openpyxl import Workbook

from threatwright.method import Method, evaluate, exposure
from threatwright.model.objects import Measure, Threat
from threatwright.model.store import Project

THREAT_COLUMNS = [
    "model", "id", "title", "scenario", "stride", "targets", "assets",
    "interface", "access", "user", "impact", "exposure", "risk", "risk_label",
    "acceptable", "status", "measures", "cra_refs", "safety_related",
]  # fmt: skip
MEASURE_COLUMNS = [
    "model", "id", "title", "kind", "status", "owner", "due", "evidence", "threats",
]  # fmt: skip
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
FIXED_ISO = "2000-01-01T00:00:00Z"
YES_NO = {"de": ("ja", "nein"), "en": ("yes", "no")}


def threat_rows(project: Project, method: Method) -> list[dict[str, Any]]:
    lang = project.config.language
    rows = []
    for name, sub in project.models.items():
        for t in sub.objects["threats"].values():
            if not isinstance(t, Threat):
                continue
            ev = evaluate(method, t, project.config)
            a = t.assessment
            label = method.risk_levels.levels.get(ev.risk) if ev.risk is not None else None
            rows.append(
                {
                    "model": name,
                    "id": t.id,
                    "title": t.title,
                    "scenario": t.scenario,
                    "stride": t.stride,
                    "targets": t.targets,
                    "assets": t.assets,
                    "interface": a.interface,
                    "access": a.access,
                    "user": a.user,
                    "impact": ev.impact,
                    "exposure": exposure(method, a),
                    "risk": ev.risk,
                    "risk_label": getattr(label, lang) if label else None,
                    "acceptable": ev.acceptable,
                    "status": t.status,
                    "measures": t.measures,
                    "cra_refs": t.cra_refs,
                    "safety_related": t.safety_related,
                }
            )
    return rows


def measure_rows(project: Project) -> list[dict[str, Any]]:
    rows = []
    for name, sub in project.models.items():
        threats = [t for t in sub.objects["threats"].values() if isinstance(t, Threat)]
        for m in sub.objects["measures"].values():
            if not isinstance(m, Measure):
                continue
            rows.append(
                {
                    "model": name,
                    "id": m.id,
                    "title": m.title,
                    "kind": m.kind,
                    "status": m.status,
                    "owner": m.owner,
                    "due": m.due.isoformat() if m.due else None,
                    "evidence": m.evidence,
                    "threats": [t.id for t in threats if m.id in t.measures],
                }
            )
    return rows


def _cell(value: Any, lang: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return YES_NO.get(lang, YES_NO["en"])[0 if value else 1]
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return str(value)


def neutralize(text: str) -> str:
    """Defuse spreadsheet formula injection (OWASP CSV injection)."""
    return "'" + text if text.startswith(FORMULA_START) else text


def _csv(rows: list[dict[str, Any]], columns: list[str], delimiter: str, lang: str) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\n")
    writer.writerow(columns)
    for row in rows:
        writer.writerow([neutralize(_cell(row[c], lang)) for c in columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _xlsx(sheets: dict[str, tuple[list[str], list[dict[str, Any]]]], lang: str) -> bytes:
    wb = Workbook()
    if wb.active is not None:
        wb.remove(wb.active)
    for title, (columns, rows) in sheets.items():
        ws = wb.create_sheet(title)
        ws.append(columns)
        for row in rows:
            values = [row[c] for c in columns]
            ws.append(
                [
                    v if isinstance(v, int | float) and not isinstance(v, bool) else _cell(v, lang)
                    for v in values
                ]
            )
        for cells in ws.iter_rows():
            for cell in cells:
                if cell.data_type == "f":  # openpyxl turns "=..." into a formula
                    cell.data_type = "s"
        ws.freeze_panes = "A2"
    raw = io.BytesIO()
    wb.save(raw)
    return _normalize_zip(raw.getvalue())


def _normalize_zip(data: bytes) -> bytes:
    """Fixed timestamps in entries and core properties: byte-identical output."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                    rb"\g<1>" + FIXED_ISO.encode(),
                    content,
                )
            entry = zipfile.ZipInfo(info.filename, FIXED_ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(entry, content)
    return out.getvalue()


def export_threats(
    project: Project,
    method: Method,
    fmt: Literal["csv", "json", "xlsx"],
    out: Path,
    delimiter: str = ";",
) -> list[Path]:
    """Write the threat list (and measure list); return the written files.

    CSV writes two files: `out` (threats) and `<stem>-measures.csv`.
    """
    lang = project.config.language
    threats, measures = threat_rows(project, method), measure_rows(project)
    out.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "csv":
        m_out = out.with_name(f"{out.stem}-measures{out.suffix or '.csv'}")
        out.write_bytes(_csv(threats, THREAT_COLUMNS, delimiter, lang))
        m_out.write_bytes(_csv(measures, MEASURE_COLUMNS, delimiter, lang))
        return [out, m_out]
    if fmt == "json":
        doc = {"threats": threats, "measures": measures}
        out.write_text(
            json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
        )
        return [out]
    if fmt == "xlsx":
        names = ("Bedrohungen", "Maßnahmen") if lang == "de" else ("Threats", "Measures")
        sheets = {
            names[0]: (THREAT_COLUMNS, threats),
            names[1]: (MEASURE_COLUMNS, measures),
        }
        out.write_bytes(_xlsx(sheets, lang))
        return [out]
    raise ValueError(f"unknown format '{fmt}'")
