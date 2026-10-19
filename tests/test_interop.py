import csv
import io
import json
import time
from pathlib import Path

import pytest
from conftest import ROOT
from openpyxl import load_workbook

from threatwright.interop import ImportFailed, export_threats, import_csv, import_threat_dragon
from threatwright.method import load_method
from threatwright.model.objects import Element, Flow, Measure, Threat
from threatwright.model.store import Project
from threatwright.rules import validate

TD_FILE = ROOT / "tests" / "fixtures" / "interop" / "threat-dragon-v2.json"


def evil_project(project_dir: Path) -> Project:
    project = Project.load(project_dir)
    project.models["main"].objects["threats"]["T-0001"].title = '=HYPERLINK("x")'
    return project


def test_csv_export(project_dir: Path, tmp_path: Path) -> None:
    project = evil_project(project_dir)
    method = load_method(project_dir, project.config)
    files = export_threats(project, method, "csv", tmp_path / "out" / "threats.csv")
    assert [f.name for f in files] == ["threats.csv", "threats-measures.csv"]
    raw = files[0].read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") and b"\r\n" not in raw
    rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    (row,) = rows
    assert row["title"] == '\'=HYPERLINK("x")'
    assert row["stride"] == "S, T, I"
    assert row["risk"] and row["risk_label"] and row["safety_related"] == "ja"
    measures = list(csv.DictReader(io.StringIO(files[1].read_text("utf-8-sig")), delimiter=";"))
    assert measures[0]["threats"] == "T-0001"
    again = export_threats(project, method, "csv", tmp_path / "b.csv", delimiter=",")
    assert again[0].read_text("utf-8-sig").splitlines()[0].startswith("model,id,title")


def test_json_export(project_dir: Path, tmp_path: Path) -> None:
    project = Project.load(project_dir)
    method = load_method(project_dir, project.config)
    (out,) = export_threats(project, method, "json", tmp_path / "t.json")
    data = json.loads(out.read_text("utf-8"))
    assert data["threats"][0]["stride"] == ["S", "T", "I"]
    assert data["threats"][0]["safety_related"] is True
    assert data["measures"][0]["id"] == "M-0001"


def test_xlsx_export_readable_safe_and_deterministic(project_dir: Path, tmp_path: Path) -> None:
    project = evil_project(project_dir)
    method = load_method(project_dir, project.config)
    (a,) = export_threats(project, method, "xlsx", tmp_path / "a.xlsx")
    time.sleep(1.1)  # zip entry and core property times would differ
    (b,) = export_threats(project, method, "xlsx", tmp_path / "b.xlsx")
    assert a.read_bytes() == b.read_bytes()
    wb = load_workbook(a)
    assert wb.sheetnames == ["Bedrohungen", "Maßnahmen"]
    ws = wb["Bedrohungen"]
    header = [c.value for c in ws[1]]
    title = ws.cell(2, header.index("title") + 1)
    assert title.value == '=HYPERLINK("x")' and title.data_type == "s"
    assert isinstance(ws.cell(2, header.index("risk") + 1).value, int)


def test_threat_dragon_import_fresh_model(project_dir: Path) -> None:
    result = import_threat_dragon(project_dir, TD_FILE, "imported")
    assert result.created_model
    assert result.counts == {"elements": 3, "zones": 1, "flows": 2, "threats": 4, "measures": 2}
    assert any("Dangling" in w for w in result.warnings)
    assert any("Network edge" in w for w in result.warnings)

    project = Project.load(project_dir)
    assert project.issues == []
    sub = project.models["imported"]
    assert sub.meta.title == "Generic demo system"
    els = {e.title: e for e in sub.objects["elements"].values() if isinstance(e, Element)}
    assert {k: (v.kind, v.zone) for k, v in els.items()} == {
        "Operator": ("external", None),
        "Controller": ("process", "Z-TB1"),
        "Parameter memory": ("datastore", "Z-TB1"),
    }
    flow = sub.objects["flows"]["F-0001"]
    assert isinstance(flow, Flow)
    assert (flow.source, flow.target, flow.bidirectional, flow.protocol) == (
        els["Operator"].id,
        els["Controller"].id,
        True,
        "HMI",
    )
    threats = {t.title: t for t in sub.objects["threats"].values() if isinstance(t, Threat)}
    fw = threats["Firmware modified"]
    assert fw.stride == ["T"] and fw.status == "identified" and fw.targets == [els["Controller"].id]
    assert fw.scenario == "Manipulated firmware is installed."
    m = sub.objects["measures"][fw.measures[0]]
    assert (
        isinstance(m, Measure)
        and m.status == "implemented"
        and m.title == "Signed firmware updates"
    )
    assert threats["Commands read on the line"].targets == ["F-0001"]
    assert threats["Commands read on the line"].stride == ["I"]
    assert "E-0001" in sub.layout.nodes

    schema = [i for i in validate(project) if i.code == "schema"]
    assert schema == []


def test_threat_dragon_import_appends_to_existing_model(project_dir: Path) -> None:
    before = (project_dir / "models" / "main" / "threats" / "T-0001.yaml").read_bytes()
    result = import_threat_dragon(project_dir, TD_FILE, "main")
    assert not result.created_model
    project = Project.load(project_dir)
    main = project.models["main"]
    assert set(main.objects["threats"]) == {"T-0001", "T-0002", "T-0003", "T-0004", "T-0005"}
    assert "E-0004" in main.objects["elements"]
    assert (project_dir / "models" / "main" / "threats" / "T-0001.yaml").read_bytes() == before
    assert [i for i in validate(project) if i.code == "schema"] == []


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"not json", "not a valid JSON"),
        (b'{"detail": {"diagrams": []}}', "no Threat Dragon diagrams"),
        (b'{"detail": {"diagrams": [{"diagramJson": {}}]}}', "v1 format"),
    ],
)
def test_threat_dragon_rejects_bad_files(
    project_dir: Path, tmp_path: Path, content: bytes, message: str
) -> None:
    f = tmp_path / "x.json"
    f.write_bytes(content)
    with pytest.raises(ImportFailed, match=message):
        import_threat_dragon(project_dir, f, "x")


def test_import_guards(project_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    with pytest.raises(ImportFailed, match="invalid model name"):
        import_threat_dragon(project_dir, TD_FILE, "../evil")
    monkeypatch.setattr("threatwright.interop.importers.MAX_IMPORT_BYTES", 10)
    with pytest.raises(ImportFailed, match="larger than"):
        import_threat_dragon(project_dir, TD_FILE, "x")


def test_csv_roundtrip(project_dir: Path, tmp_path: Path) -> None:
    project = evil_project(project_dir)
    project.save()
    method = load_method(project_dir, project.config)
    out, _ = export_threats(project, method, "csv", tmp_path / "t.csv")
    result = import_csv(project_dir, out, "main")
    assert result.counts == {"threats": 1}
    t = Project.load(project_dir).models["main"].objects["threats"]["T-0002"]
    assert isinstance(t, Threat)
    assert t.title == '=HYPERLINK("x")'
    assert t.stride == ["S", "T", "I"]
    assert t.targets == ["E-0002", "F-0001"]
    assert t.measures == ["M-0001"] and t.safety_related is True
    assert t.assessment.impact == 4 and t.assessment.interface == "local"


def test_csv_import_comma_and_warnings(project_dir: Path, tmp_path: Path) -> None:
    f = tmp_path / "in.csv"
    f.write_text(
        "title,stride,targets,status,impact\n"
        'Plain threat,"Spoofing, Tampering",E-0001,rejected,9\n'
        "Other,X,E-9999,assessed,\n"
        ",,,,\n",
        encoding="utf-8",
    )
    result = import_csv(project_dir, f, "fresh")
    assert result.created_model and result.counts == {"threats": 1}
    text = "\n".join(result.warnings)
    assert "unknown STRIDE category 'X'" in text
    assert "E-9999" in text and "E-0001" in text  # fresh model has no elements
    assert "line 2: skipped" in text  # impact 9 out of range
    sub = Project.load(project_dir).models["fresh"]
    (t,) = sub.objects["threats"].values()
    assert isinstance(t, Threat) and t.title == "Other" and t.status == "assessed"


def test_csv_import_requires_title_header(project_dir: Path, tmp_path: Path) -> None:
    f = tmp_path / "in.csv"
    f.write_text("a;b\n1;2\n", encoding="utf-8")
    with pytest.raises(ImportFailed, match="title"):
        import_csv(project_dir, f, "main")
