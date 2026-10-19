import zipfile
from pathlib import Path

import docx
import pytest
from lxml import etree

from threatwright.method import load_method
from threatwright.model.store import Project
from threatwright.render import (
    KNOWN_VARIABLES,
    RenderError,
    build_context,
    build_documents,
    check_template,
    derive_template,
    find_soffice,
    list_templates,
    preview_html,
)
from threatwright.render import pdf as pdf_module
from threatwright.render.package import BUILTIN_TEMPLATES, available

BUNDLED = ["acceptance-criteria", "tr-analysis", "tr-analysis-compact"]


def text_of(path: Path) -> str:
    if path.suffix in (".md", ".html"):
        return path.read_text(encoding="utf-8")
    if path.suffix == ".docx":
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
        parts += [p.text for s in d.sections for p in s.header.paragraphs + s.footer.paragraphs]
        return "\n".join(parts)
    with zipfile.ZipFile(path) as zf:
        roots = [etree.fromstring(zf.read(n)) for n in ("content.xml", "styles.xml")]
    return "\n".join("".join(r.itertext()) for r in roots)


def headings(path: Path) -> list[str]:
    return [p.text for p in docx.Document(str(path)).paragraphs if p.style.name == "Heading 1"]


def set_language(project_dir: Path, lang: str) -> None:
    cfg = project_dir / "threatwright.yaml"
    cfg.write_text(
        cfg.read_text(encoding="utf-8").replace("language: de", f"language: {lang}"),
        encoding="utf-8",
    )


def test_build_all_bundled_templates(project_dir: Path) -> None:
    paths = build_documents(project_dir)
    names = sorted(p.name for p in paths)
    assert names == sorted(
        [
            "acceptance-criteria.docx",
            "acceptance-criteria.md",
            "acceptance-criteria.odt",
            "main-tr-analysis-compact.docx",
            "main-tr-analysis.docx",
            "main-tr-analysis.html",
            "main-tr-analysis.md",
            "main-tr-analysis.odt",
        ]
    )
    assert all(p.parent == project_dir / "build" for p in paths)
    for p in paths:
        text = text_of(p)
        assert "Test fixture project" in text, p.name
        assert "{{" not in text and "{%" not in text, p.name
        if p.stem == "acceptance-criteria":
            for needle in (
                "Höchste akzeptable Risikostufe",
                "Unterschrift",
                "umgekehrt (inverted)",
                "Zu behandelnde Bedrohungen",
                "Produktverantwortung",
            ):
                assert needle in text, (p.name, needle)
        else:
            for needle in ("T-0001", "3 – mittel", "zu behandeln", "M-0001"):
                assert needle in text, (p.name, needle)
        if p.stem == "main-tr-analysis":
            # method section: visible matrix-reading note and acceptance limit
            assert "Lesart der Risikomatrix: umgekehrt (inverted)" in text
            assert "Akzeptanzgrenze: Risiken bis einschließlich „niedrig“" in text
            assert "Safety functions" in text and "Functions.Safety" in text
            assert "Q-0001" in text and "AI-I-2-e" in text


def test_english_and_literal_reading(project_dir: Path) -> None:
    set_language(project_dir, "en")
    (project_dir / "method-overrides.yaml").write_text(
        "risk_matrix:\n  matrix_row_from_exposure: literal\n", encoding="utf-8"
    )
    [md] = build_documents(project_dir, template="tr-analysis", fmt="md")
    text = md.read_text(encoding="utf-8")
    assert "Risk matrix reading: literal" in text
    assert "4 – high" in text and "to be treated" in text  # literal row 2, impact 4


def test_rejected_threats_have_own_section(project_dir: Path) -> None:
    threat = project_dir / "models" / "main" / "threats" / "T-0001.yaml"
    threat.write_text(
        threat.read_text(encoding="utf-8").replace("status: assessed", "status: rejected")
        + "rejection:\n  by: Reviewer\n  rationale: Not applicable here.\n",
        encoding="utf-8",
    )
    [md] = build_documents(project_dir, template="tr-analysis", fmt="md")
    text = md.read_text(encoding="utf-8")
    rejected = text.split("## Verworfene Bedrohungen")[1].split("## ")[0]
    assert "T-0001" in rejected and "Not applicable here." in rejected
    distribution = text.split("## Risikoverteilung")[1].split("## ")[0]
    assert "| 3 | mittel | 0 |" in distribution


def test_acceptance_appendix_counts(project_dir: Path) -> None:
    project = Project.load(project_dir)
    method = load_method(project_dir, project.config)
    meta = available(project_dir)["acceptance-criteria"].meta
    ctx = build_context(project, method, meta, None)
    decisions = {d["id"]: d for d in ctx["acceptance"]["decisions"]}
    limit = decisions["max_acceptable_risk"]["options"]
    assert [o["to_treat"] for o in limit] == [1, 1, 0, 0, 0]
    assert [o["recommended"] for o in limit] == [False, True, False, False, False]
    reading = {o["label"]: o["to_treat"] for o in decisions["matrix_reading"]["options"]}
    assert reading == {"umgekehrt (inverted)": 1, "wörtlich (literal)": 1}
    post_market = decisions["rule:post-market-moderate"]["options"]
    assert [o["to_treat"] for o in post_market] == [1, 1]  # lifecycle is development
    assert set(ctx) == KNOWN_VARIABLES


def test_md_html_docx_odt_are_deterministic(project_dir: Path, tmp_path: Path) -> None:
    first = build_documents(project_dir, template="tr-analysis", out_dir=tmp_path / "a")
    second = build_documents(project_dir, template="tr-analysis", out_dir=tmp_path / "b")
    assert len(first) == 4
    for a, b in zip(first, second, strict=True):
        assert a.read_bytes() == b.read_bytes(), a.name


def test_compact_template_changes_order_and_columns(project_dir: Path) -> None:
    std = build_documents(project_dir, template="tr-analysis", fmt="docx")[0]
    compact = build_documents(project_dir, template="tr-analysis-compact", fmt="docx")[0]
    h_std, h_compact = headings(std), headings(compact)
    assert h_std.index("Methode") < h_std.index("Risikoverteilung")
    assert h_compact.index("Risikoverteilung") < h_compact.index("Methode")
    assert h_compact[0] == "Zusammenfassung" and "Assets" not in h_compact
    threat_tables = [
        t for t in docx.Document(str(compact)).tables if t.rows[0].cells[0].text == "ID"
    ]
    assert [c.text for c in threat_tables[0].rows[0].cells] == [
        "ID",
        "Titel",
        "Risiko",
        "Bewertung",
    ]
    styles = {s.name: s for s in docx.Document(str(compact)).styles}
    assert styles["Normal"].font.name == "Arial"


def _project_template(project_dir: Path, tid: str, body: str) -> Path:
    pkg = project_dir / "templates" / tid
    pkg.mkdir(parents=True)
    (pkg / "template.yaml").write_text(
        f"schema_version: 1\nid: {tid}\nname: {{de: X, en: X}}\nformat: md\nfile: t.md\n"
        "outline:\n  - type: cover\n",
        encoding="utf-8",
    )
    (pkg / "t.md").write_text(body, encoding="utf-8")
    return pkg


@pytest.mark.parametrize(
    "body",
    [
        "{{ ''.__class__.__mro__[1].__subclasses__() }}",
        "{{ project.__class__ }}{{ project.__class__.__init__ }}",
        "{% include '../../threatwright.yaml' %}",
        "{{ threats.append(1) }}",
    ],
)
def test_templates_run_sandboxed(project_dir: Path, body: str) -> None:
    _project_template(project_dir, "evil", body)
    with pytest.raises(RenderError):
        build_documents(project_dir, template="evil")


def test_free_text_runs_sandboxed(project_dir: Path) -> None:
    pkg = _project_template(project_dir, "evil-text", "{{ outline[0].blocks[0].text }}")
    (pkg / "template.yaml").write_text(
        "schema_version: 1\nid: evil-text\nname: {de: X, en: X}\nformat: md\nfile: t.md\n"
        "outline:\n  - type: text\n    text: \"{{ ''.__class__.__mro__ }}\"\n",
        encoding="utf-8",
    )
    with pytest.raises(RenderError, match="template.yaml"):
        build_documents(project_dir, template="evil-text")


def test_docx_templates_run_sandboxed(project_dir: Path) -> None:
    pkg = project_dir / "templates" / "evil-docx"
    pkg.mkdir(parents=True)
    d = docx.Document()
    d.add_paragraph("{{ ''.__class__.__mro__[1].__subclasses__() }}")
    d.save(str(pkg / "t.docx"))
    (pkg / "template.yaml").write_text(
        "schema_version: 1\nid: evil-docx\nname: {de: X, en: X}\nformat: docx\nfile: t.docx\n"
        "outline:\n  - type: cover\n",
        encoding="utf-8",
    )
    with pytest.raises(RenderError):
        build_documents(project_dir, template="evil-docx")


def test_project_template_overrides_bundled(project_dir: Path) -> None:
    pkg = _project_template(project_dir, "tr-analysis", "OWN {{ project.name }}\n")
    rows = {r["id"]: r for r in list_templates(project_dir)}
    assert rows["tr-analysis"]["origin"] == "project" and rows["tr-analysis"]["formats"] == ["md"]
    assert rows["acceptance-criteria"]["origin"] == "bundled"
    assert all(rows[i]["active"] for i in BUNDLED)
    [md] = build_documents(project_dir, template="tr-analysis")
    assert md.read_text(encoding="utf-8") == "OWN Test fixture project\n"
    assert pkg.exists()


def test_project_config_selects_templates(project_dir: Path) -> None:
    cfg = project_dir / "threatwright.yaml"
    cfg.write_text(
        cfg.read_text(encoding="utf-8") + "templates: [acceptance-criteria]\n", encoding="utf-8"
    )
    assert sorted(p.name for p in build_documents(project_dir)) == [
        "acceptance-criteria.docx",
        "acceptance-criteria.md",
        "acceptance-criteria.odt",
    ]
    assert not {r["id"]: r for r in list_templates(project_dir)}["tr-analysis"]["active"]


def test_check_bundled_templates_clean() -> None:
    for tid in BUNDLED:
        assert check_template(BUILTIN_TEMPLATES / tid) == [], tid


def test_check_finds_unknown_variable_and_syntax_error(tmp_path: Path) -> None:
    pkg = _project_template(tmp_path, "bad", "{{ project.name }} {{ threatz }}\n")
    issues = check_template(pkg)
    assert [(i.code, i.message) for i in issues] == [
        ("template-variable", "unknown variable 'threatz'")
    ]
    (pkg / "t.md").write_text("{% for x in threats %}{{ x.id }}", encoding="utf-8")
    assert [i.code for i in check_template(pkg / "template.yaml")] == ["template-syntax"]
    (pkg / "t.md").write_text("{{ threats|no_such_filter }}", encoding="utf-8")
    assert [i.code for i in check_template(pkg)] == ["template-syntax"]
    (pkg / "template.yaml").write_text("schema_version: 1\nid: bad\n", encoding="utf-8")
    assert [i.code for i in check_template(pkg)] == ["template-yaml"]


def test_check_docx_unknown_variable(tmp_path: Path) -> None:
    pkg = tmp_path / "templates" / "w"
    pkg.mkdir(parents=True)
    d = docx.Document()
    d.add_paragraph("{{ document.title }} {{ unknown_thing }}")
    d.save(str(pkg / "t.docx"))
    (pkg / "template.yaml").write_text(
        "schema_version: 1\nid: w\nname: {de: W, en: W}\nformat: docx\nfile: t.docx\n"
        "outline:\n  - type: text\n    text: '{{ nope }}'\n",
        encoding="utf-8",
    )
    messages = sorted(i.message for i in check_template(pkg))
    assert messages == ["unknown variable 'nope'", "unknown variable 'unknown_thing'"]


def test_template_yaml_rejects_paths_and_bad_columns(tmp_path: Path) -> None:
    pkg = _project_template(tmp_path, "p", "x")
    (pkg / "template.yaml").write_text(
        "schema_version: 1\nid: p\nname: {de: P, en: P}\nformat: md\nfile: ../t.md\n"
        "outline:\n  - type: cover\n",
        encoding="utf-8",
    )
    assert "plain file name" in check_template(pkg)[0].message
    (pkg / "template.yaml").write_text(
        "schema_version: 1\nid: p\nname: {de: P, en: P}\nformat: md\nfile: t.md\n"
        "outline:\n  - type: threats\n    columns: [id, colour]\n",
        encoding="utf-8",
    )
    assert "unknown column" in check_template(pkg)[0].message


def test_derive_docx_roundtrip(project_dir: Path, tmp_path: Path) -> None:
    source = tmp_path / "Hersteller Vorlage.docx"
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Kopfzeile des Herstellers"
    d.add_heading("Altes Kapitel", 1)
    d.add_paragraph("Alter Inhalt, der ersetzt wird.")
    d.save(str(source))
    target = project_dir / "templates" / "hersteller"
    meta = derive_template(source, target)
    assert meta == target / "template.yaml"
    assert check_template(target) == []
    [out] = build_documents(project_dir, template="hersteller")
    text = text_of(out)
    assert "Kopfzeile des Herstellers" in text
    assert "Alter Inhalt" not in text and "Altes Kapitel" not in text
    assert "T-0001" in text and "3 – mittel" in text
    with pytest.raises(RenderError):
        derive_template(source, target)  # never overwrites


def test_derive_odt_roundtrip(project_dir: Path, tmp_path: Path) -> None:
    [built] = build_documents(project_dir, template="tr-analysis", fmt="odt", out_dir=tmp_path)
    target = project_dir / "templates" / "odt-copy"
    derive_template(built, target)
    assert check_template(target) == []
    [out] = build_documents(project_dir, template="odt-copy")
    text = text_of(out)
    assert text.count("Lesart der Risikomatrix: umgekehrt") == 1  # old body removed
    assert "{{" not in text


def test_preview_html_is_self_contained(project_dir: Path) -> None:
    html = preview_html(project_dir, "main", "tr-analysis")
    assert html.startswith("<!DOCTYPE html>") and "T-0001" in html
    for needle in ("<script", "<link", "src=", "url("):
        assert needle not in html
    # templates without an HTML file fall back to the generic layout
    assert "Höchste akzeptable Risikostufe" in preview_html(
        project_dir, None, "acceptance-criteria"
    )
    with pytest.raises(RenderError):
        preview_html(project_dir, "nope", "tr-analysis")


def test_html_escapes_model_text(project_dir: Path) -> None:
    threat = project_dir / "models" / "main" / "threats" / "T-0001.yaml"
    threat.write_text(
        threat.read_text(encoding="utf-8").replace(
            "consequence: Manipulated inputs.", "consequence: '<b>x</b> {{ 7*7 }}'"
        ),
        encoding="utf-8",
    )
    html = preview_html(project_dir, "main", "tr-analysis")
    assert "&lt;b&gt;x&lt;/b&gt; {{ 7*7 }}" in html
    [odt] = build_documents(project_dir, template="tr-analysis", fmt="odt")
    assert "{{ 7*7 }}" in text_of(odt)  # model data is never evaluated as a template


def test_list_templates_bundled_only() -> None:
    rows = {r["id"]: r for r in list_templates()}
    assert sorted(rows) == BUNDLED
    assert rows["tr-analysis"]["formats"] == ["docx", "odt", "md", "html", "pdf"]
    assert rows["acceptance-criteria"]["scope"] == "project"


def test_pdf_without_libreoffice(project_dir: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(pdf_module, "find_soffice", lambda: None)
    with pytest.raises(RenderError, match="LibreOffice"):
        build_documents(project_dir, template="tr-analysis", fmt="pdf")
    with pytest.raises(RenderError, match="cannot produce pdf"):
        _project_template(project_dir, "mdonly", "x")
        build_documents(project_dir, template="mdonly", fmt="pdf")


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_pdf_with_libreoffice(project_dir: Path) -> None:
    [out] = build_documents(project_dir, template="tr-analysis-compact", fmt="pdf")
    assert out.name == "main-tr-analysis-compact.pdf"
    assert out.read_bytes().startswith(b"%PDF")


def test_unknown_model_and_template(project_dir: Path) -> None:
    with pytest.raises(RenderError, match="model"):
        build_documents(project_dir, model="nope")
    with pytest.raises(RenderError, match="template"):
        build_documents(project_dir, template="nope")
